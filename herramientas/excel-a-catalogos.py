"""Convierte "LISTA DESPLEGABLE APP COMEX.xlsx" en catalogos.js (las listas desplegables de la app).

Uso:  python herramientas/excel-a-catalogos.py "LISTA DESPLEGABLE APP COMEX.xlsx" catalogos.js

Estructura esperada del Excel (la que entregó Ventas):
  Hoja1: columnas Cliente | (vacía) | Código | Producto        → productos por cliente
  Hoja2: columnas Destino | | Lugar de Carga | | Naviera        → listas sueltas
Las columnas se buscan por el título de la primera fila, así que pueden moverse de lugar.
"""
import json
import sys
from datetime import date

import openpyxl


def columnas(ws):
    titulos = [str(c.value or '').strip().lower() for c in ws[1]]
    return {t: i for i, t in enumerate(titulos) if t}


def texto(v):
    return ' '.join(str(v).split()) if v is not None else ''


def main(origen, destino):
    wb = openpyxl.load_workbook(origen, data_only=True)
    productos, listas = [], {'destinos': [], 'lugaresCarga': [], 'navieras': [], 'transportistas': []}
    nombres = {'destino': 'destinos', 'lugar de carga': 'lugaresCarga', 'naviera': 'navieras', 'transportista': 'transportistas'}
    for ws in wb.worksheets:
        col = columnas(ws)
        filas = list(ws.iter_rows(min_row=2, values_only=True))
        if {'cliente', 'producto'} <= col.keys():
            for f in filas:
                cli, prod = texto(f[col['cliente']]), texto(f[col['producto']])
                cod = texto(f[col['código']]) if 'código' in col else texto(f[col['codigo']]) if 'codigo' in col else ''
                if cli and prod:
                    productos.append({'cliente': cli, 'producto': prod, 'codigo': cod})
        for titulo, clave in nombres.items():
            if titulo in col:
                for f in filas:
                    v = texto(f[col[titulo]])
                    if v and v not in listas[clave]:
                        listas[clave].append(v)
    clientes = sorted({p['cliente'] for p in productos})
    productos.sort(key=lambda p: (p['cliente'], p['producto'].lower()))
    datos = {'clientes': clientes, **listas, 'productos': productos}
    with open(destino, 'w', encoding='utf-8') as f:
        f.write(f'// Listas desplegables de Hub Comex. Generado desde "{origen.split(chr(92))[-1]}" el {date.today():%d/%m/%Y}.\n')
        f.write('// Para actualizar: cambiar el Excel y volver a ejecutar herramientas/excel-a-catalogos.py\n')
        f.write('window.HUB_CATALOGOS = ')
        f.write(json.dumps(datos, ensure_ascii=False, indent=1))
        f.write(';\n')
    print(f'{len(clientes)} clientes, {len(productos)} productos, ' + ', '.join(f'{len(v)} {k}' for k, v in listas.items()))


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
